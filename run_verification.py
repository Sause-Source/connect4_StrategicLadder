#!/usr/bin/env python
"""Connect4 solver: depth-vs-depth tournament -> Excel report.

Asks for the board size (width x height) and the minimum/maximum depth, plays
every (first depth, second depth) combination with connect4/verify_<w>x<h>.exe
and writes an .xlsx file containing

  * 対戦結果 : matrix (rows = first player, columns = second player)
               red = first player wins, blue = second player wins, grey = draw
  * 勝率     : per-depth "段" (tier) ranking + win rate against the current top
               tier among lower depths, and a line chart of that win rate

Tier ("段") ranking: depth `lo` starts at tier 0. For each following depth d
(in increasing order), let T = the highest tier reached by any depth < d, and
compute d's win rate against the depths that are AT tier T. If that win rate is
above LEVEL_UP_THRESHOLD, d reaches tier T+1; otherwise d stays at tier T.

The board size is a compile-time constant of the solver (bitboard packing), so
a new size is built once (as connect4/verify_<w>x<h>.exe) and reused afterwards.
Game results are accumulated per board size in results_<w>x<h>.csv, so an
interrupted run (or a later run with a wider depth range) only plays the games
that are still missing.

usage: python run_verification.py [min_depth max_depth] [--width W] [--height H]
                                   [--threads N] [--report-only]
"""
import argparse
import csv
import os
import shutil
import subprocess
import sys

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.abspath(__file__))
SRC_DIR = os.path.join(ROOT, "connect4")
SOURCES = ["verify.cpp", "Solver.cpp", "Solver.hpp", "Position.hpp", "MoveSorter.hpp",
           "TranspositionTable.hpp", "OpeningBook.hpp"]
STANDARD_BOOK = os.path.join(SRC_DIR, "7x6.book")  # opening book only covers the standard 7x6 board

# Solver constraints (see the static_asserts in connect4/Position.hpp): the column
# index is packed on a byte (< 10 columns) and the whole board must fit in 128 bits.
MIN_SIZE = 1
MAX_WIDTH = 9

FIRST_WIN, SECOND_WIN, DRAW = 1, 2, 0
RED, BLUE, GREY = "E15759", "4E79A7", "D9D9D9"

# Win rate (against the current top tier) above which a depth levels up one tier.
# Adjust this value to make leveling up stricter (closer to 1) or easier (closer to 0).
LEVEL_UP_THRESHOLD = 0.6


def cells(width, height):
    return width * height


def exe_ext():
    return ".exe" if os.name == "nt" else ""


def exe_path(width, height):
    return os.path.join(SRC_DIR, "verify_%dx%d%s" % (width, height, exe_ext()))


def results_path(width, height):
    return os.path.join(ROOT, "results_%dx%d.csv" % (width, height))


def book_path(width, height):
    if (width, height) == (7, 6) and os.path.exists(STANDARD_BOOK):
        return STANDARD_BOOK
    return None  # no book for non-standard board sizes


def validate_board(width, height):
    if not (MIN_SIZE <= width <= MAX_WIDTH):
        sys.exit("列数(width)は %d〜%d の整数にしてください(ソルバーの制約)。" % (MIN_SIZE, MAX_WIDTH))
    if height < MIN_SIZE:
        sys.exit("行数(height)は %d 以上の整数にしてください。" % MIN_SIZE)
    if width * (height + 1) > 128:
        sys.exit("列数 x (行数+1) = %d が 128 を超えています。もっと小さい盤面にしてください。"
                  % (width * (height + 1)))


# ---------------------------------------------------------------- build / run
def find_compiler():
    """Returns a command prefix for a C++ compiler, or None."""
    for name in ("g++", "clang++"):
        if shutil.which(name):
            return [shutil.which(name)]
    zig = os.path.join(ROOT, ".tools", "Lib", "site-packages", "ziglang", "zig.exe")
    if os.path.exists(zig):
        return [zig, "c++"]
    return None


def build(width, height):
    exe = exe_path(width, height)
    newest = max(os.path.getmtime(os.path.join(SRC_DIR, s)) for s in SOURCES)
    if os.path.exists(exe) and os.path.getmtime(exe) >= newest:
        return exe
    cc = find_compiler()
    if cc is None:
        sys.exit("C++ compiler (g++/clang++) not found and %s is missing." % exe)
    print("building %s (board %dx%d) ..." % (exe, width, height))
    cmd = cc + ["-std=c++17", "-O3", "-DNDEBUG",
                "-DC4_WIDTH=%d" % width, "-DC4_HEIGHT=%d" % height,
                "-o", exe, os.path.join(SRC_DIR, "verify.cpp"), os.path.join(SRC_DIR, "Solver.cpp")]
    subprocess.run(cmd, check=True, cwd=SRC_DIR)
    return exe


def run_games(width, height, lo, hi, threads):
    book = book_path(width, height)
    if book is None:
        print("note: no opening book for a %dx%d board; solving will be slower (perfect play near the end of the game only)."
              % (width, height))
    subprocess.run([exe_path(width, height), str(lo), str(hi), results_path(width, height),
                    str(threads), book or "-"], check=True, cwd=SRC_DIR)


# -------------------------------------------------------------------- report
def load_results(width, height):
    res = {}
    path = results_path(width, height)
    if not os.path.exists(path):
        return res
    with open(path, newline="") as f:
        for row in csv.DictReader(f):
            res[(int(row["first"]), int(row["second"]))] = (int(row["result"]), int(row["moves"]))
    return res


def games_against(res, d, group):
    """Aggregate d's games (as first and as second player) against every depth in group."""
    s = dict(games=0, win=0, draw=0, loss=0, f_games=0, f_win=0, s_games=0, s_win=0)
    for e in group:
        if (d, e) in res:  # d plays first
            r = res[(d, e)][0]
            s["games"] += 1; s["f_games"] += 1
            if r == FIRST_WIN: s["win"] += 1; s["f_win"] += 1
            elif r == DRAW: s["draw"] += 1
            else: s["loss"] += 1
        if (e, d) in res:  # d plays second
            r = res[(e, d)][0]
            s["games"] += 1; s["s_games"] += 1
            if r == SECOND_WIN: s["win"] += 1; s["s_win"] += 1
            elif r == DRAW: s["draw"] += 1
            else: s["loss"] += 1
    return s


def compute_tiers(res, depths, threshold=LEVEL_UP_THRESHOLD):
    """Assigns each depth a tier ("段"), starting at 0 for the lowest depth.

    For each depth d (processed in increasing order, so every lower depth's tier
    is already known), the "top tier" T is the highest tier reached by any depth
    below d. d's win rate is computed only against the depths that are at tier T
    (games as first and as second player, draws count as neither win nor loss).
    If that win rate is strictly above `threshold`, d reaches tier T+1; otherwise
    it stays at T.

    Returns (tier, info) where tier: depth -> int, and info: depth -> dict with
    the reference tier used, the game counts/rates against it, and whether the
    depth leveled up.
    """
    tier = {}
    info = {}
    for i, d in enumerate(depths):
        lower = depths[:i]
        if not lower:
            tier[d] = 0
            info[d] = dict(ref_tier=None, games=0, win=0, draw=0, loss=0, f_games=0, f_win=0,
                            s_games=0, s_win=0, winrate=None, f_rate=None, s_rate=None, leveled_up=False)
            continue
        top_tier = max(tier[e] for e in lower)
        group = [e for e in lower if tier[e] == top_tier]
        s = games_against(res, d, group)
        winrate = rate(s["win"], s["games"])
        leveled_up = winrate is not None and winrate > threshold
        tier[d] = top_tier + 1 if leveled_up else top_tier
        info[d] = dict(ref_tier=top_tier, **s, winrate=winrate,
                        f_rate=rate(s["f_win"], s["f_games"]), s_rate=rate(s["s_win"], s["s_games"]),
                        leveled_up=leveled_up)
    return tier, info


def rate(a, b):
    return a / b if b else None


def write_report(res, width, height, lo, hi, path):
    depths = list(range(lo, hi + 1))
    wb = Workbook()
    thin = Side(style="thin", color="FFFFFF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")
    bold_white = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="404040")

    # ---- sheet 1: matrix
    ws = wb.active
    ws.title = "対戦結果(%dx%d)" % (width, height)
    ws["A1"] = "先手＼後手"
    ws["A1"].font = bold_white; ws["A1"].fill = head_fill; ws["A1"].alignment = center
    ws.column_dimensions["A"].width = 12
    for i, d in enumerate(depths):
        for cell in (ws.cell(row=1, column=i + 2, value=d), ws.cell(row=i + 2, column=1, value=d)):
            cell.font = bold_white; cell.fill = head_fill; cell.alignment = center
        ws.column_dimensions[get_column_letter(i + 2)].width = 4.5
    fills = {FIRST_WIN: PatternFill("solid", fgColor=RED),
             SECOND_WIN: PatternFill("solid", fgColor=BLUE),
             DRAW: PatternFill("solid", fgColor=GREY)}
    for i, a in enumerate(depths):
        for j, b in enumerate(depths):
            if (a, b) not in res:
                continue
            r, moves = res[(a, b)]
            c = ws.cell(row=i + 2, column=j + 2, value=moves)  # number of moves in the game
            c.fill = fills[r]; c.border = border; c.alignment = center
            c.font = Font(color="FFFFFF" if r != DRAW else "000000", size=9)
    ws.freeze_panes = "B2"
    lg = len(depths) + 4
    ws.cell(row=lg, column=1, value="凡例").font = Font(bold=True)
    for k, (label, colour) in enumerate([("先手(行)の勝ち", RED), ("後手(列)の勝ち", BLUE), ("引き分け", GREY)]):
        ws.cell(row=lg + 1 + k, column=1).fill = PatternFill("solid", fgColor=colour)
        ws.cell(row=lg + 1 + k, column=2, value=label)
    ws.cell(row=lg + 4, column=2, value="セル内の数字 = その対局の総手数")
    ws.cell(row=lg + 5, column=2,
            value="盤面サイズ: %d列 x %d行(総マス数 %d、最大深さ=総マス数)" % (width, height, cells(width, height)))

    # ---- sheet 2: tiers ("段") and win rate against the current top tier
    wr = wb.create_sheet("勝率(%dx%d)" % (width, height))
    headers = ["深さ", "比較対象の段", "対戦数", "勝", "分", "負", "勝率", "先手時の勝率", "後手時の勝率", "段"]
    COL_RATE, COL_TIER = 7, 10  # 1-based column numbers used by the chart / level-up highlight below
    for k, h in enumerate(headers):
        c = wr.cell(row=1, column=k + 1, value=h)
        c.font = bold_white; c.fill = head_fill; c.alignment = center
        wr.column_dimensions[get_column_letter(k + 1)].width = 14
    tier, info = compute_tiers(res, depths)
    level_up_fill = PatternFill("solid", fgColor="C6E0B4")  # highlights the depths that leveled up
    for i, d in enumerate(depths):
        s = info[d]
        row = [d, s["ref_tier"], s["games"], s["win"], s["draw"], s["loss"],
               s["winrate"], s["f_rate"], s["s_rate"], tier[d]]
        for k, v in enumerate(row):
            c = wr.cell(row=i + 2, column=k + 1, value=v if v is not None else "-")
            c.alignment = center
            if k + 1 in (COL_RATE, 8, 9) and v is not None:
                c.number_format = "0.0%"
        if s["leveled_up"]:
            wr.cell(row=i + 2, column=COL_TIER).fill = level_up_fill
    note_row = len(depths) + 3
    wr.cell(row=note_row, column=1,
            value="勝率 = 自分未満の深さの中で最大の段に属する深さに対する勝ち数 / 対戦数"
                  "(先手・後手の両方を含む。引き分けは勝ちに含めない)")
    wr.cell(row=note_row + 1, column=1,
            value="段 = 上記の勝率が %.0f%% を超えたら比較対象の段より1つ上がる、そうでなければ同じ段のまま"
                  "(先頭の深さ %d は初期値として段0)" % (LEVEL_UP_THRESHOLD * 100, depths[0]))
    wr.cell(row=note_row + 2, column=1,
            value="段の初期値は0。しきい値(現在 %.0f%%)は run_verification.py の LEVEL_UP_THRESHOLD で変更可能"
                  % (LEVEL_UP_THRESHOLD * 100))
    wr.cell(row=note_row + 3, column=1,
            value="到達した最大の段: %d(緑色のセルがその段に上がった深さ)" % max(tier.values()))

    chart = LineChart()
    chart.title = "深さごとの勝率(自分未満の深さの中で最大の段に対して) 盤面 %dx%d" % (width, height)
    chart.x_axis.title = "深さ"
    chart.y_axis.title = "勝率"
    chart.y_axis.scaling.min = 0
    chart.y_axis.scaling.max = 1
    chart.y_axis.number_format = "0%"
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.height, chart.width = 10, max(18, min(60, len(depths) * 0.9))
    data = Reference(wr, min_col=COL_RATE, min_row=1, max_row=len(depths) + 1)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(Reference(wr, min_col=1, min_row=2, max_row=len(depths) + 1))
    chart.series[0].marker.symbol = "circle"
    chart.legend = None
    wr.add_chart(chart, "L2")

    tier_chart = LineChart()
    tier_chart.title = "深さごとの段 盤面 %dx%d" % (width, height)
    tier_chart.x_axis.title = "深さ"
    tier_chart.y_axis.title = "段"
    tier_chart.x_axis.delete = False
    tier_chart.y_axis.delete = False
    tier_chart.height, tier_chart.width = 10, max(18, min(60, len(depths) * 0.9))
    tier_data = Reference(wr, min_col=COL_TIER, min_row=1, max_row=len(depths) + 1)
    tier_chart.add_data(tier_data, titles_from_data=True)
    tier_chart.set_categories(Reference(wr, min_col=1, min_row=2, max_row=len(depths) + 1))
    tier_chart.series[0].marker.symbol = "circle"
    tier_chart.legend = None
    wr.add_chart(tier_chart, "L22")

    wb.save(path)


# ---------------------------------------------------------------------- main
def ask_int(prompt, default, lo, hi):
    while True:
        text = input("%s [%s]: " % (prompt, default)).strip()
        if not text:
            return default
        if text.isdigit() and lo <= int(text) <= hi:
            return int(text)
        print("%d〜%d の整数を入力してください。" % (lo, hi))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("depths", nargs="*", type=int, help="min_depth max_depth")
    ap.add_argument("--width", type=int, help="盤面の列数 (default: 7)")
    ap.add_argument("--height", type=int, help="盤面の行数 (default: 6)")
    ap.add_argument("--threads", type=int, default=0, help="worker threads (default: min(8, cores))")
    ap.add_argument("--report-only", action="store_true", help="do not play, only build the report from results_<w>x<h>.csv")
    args = ap.parse_args()

    if args.width is not None and args.height is not None:
        width, height = args.width, args.height
    else:
        width = ask_int("盤面の列数(width)", args.width or 7, MIN_SIZE, MAX_WIDTH)
        height = ask_int("盤面の行数(height)", args.height or 6, MIN_SIZE, 128 // width - 1)
    validate_board(width, height)
    max_depth = cells(width, height)

    if len(args.depths) == 2:
        lo, hi = args.depths
    else:
        lo = ask_int("最小深さ", 1, 1, max_depth)
        hi = ask_int("最大深さ", min(10, max_depth), 1, max_depth)
    if not (1 <= lo <= hi <= max_depth):
        sys.exit("depths must satisfy 1 <= min <= max <= %d (board has %d cells)" % (max_depth, max_depth))

    if not args.report_only:
        build(width, height)
        threads = args.threads or min(8, os.cpu_count() or 1)
        run_games(width, height, lo, hi, threads)

    res = load_results(width, height)
    out = os.path.join(ROOT, "connect4_result_%dx%d_%d-%d.xlsx" % (width, height, lo, hi))
    write_report(res, width, height, lo, hi, out)
    played = sum(1 for a in range(lo, hi + 1) for b in range(lo, hi + 1) if (a, b) in res)
    print("games recorded: %d / %d" % (played, (hi - lo + 1) ** 2))
    print("report written:", out)


if __name__ == "__main__":
    main()
